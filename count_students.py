#!/usr/bin/env python3
"""Count students in the Data Nilai Excel file."""
import openpyxl

wb = openpyxl.load_workbook(
    "/opt/data/home/workspace/dev/pondok-banat/.data-pondokanat/Data Nilai Semester 1 TA 2026_2027.xlsx",
    read_only=True,
)
total = 0
for name in wb.sheetnames:
    ws = wb[name]
    cnt = 0
    for row in ws.iter_rows(values_only=True):
        if row[0] is not None and isinstance(row[0], (int, float)) and row[0] > 0:
            cnt += 1
    total += cnt
    print(f"{name}: {cnt} students")
print(f"TOTAL: {total} students")
wb.close()
