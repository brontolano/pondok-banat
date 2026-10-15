import openpyxl
from collections import Counter

# Verify teacher count in DATA PARA PENGAJAR BANAT.xlsx
f = '.data-pondokbanat/DATA PARA PENGAJAR BANAT.xlsx'
wb = openpyxl.load_workbook(f, read_only=True)

# Table 1
ws = wb['Table 1']
rows = list(ws.iter_rows(values_only=True))
teachers = []
for r in rows[2:]:
    if r[0] is not None and r[1]:
        teachers.append((r[0], r[1].strip(), r[5]))
print('=== Table 1: teacher records (NO, NAMA, JABATAN):', len(teachers))
for t in teachers[:5]:
    print('  ', t)
print('  ...')
# Check duplicates
names = [t[1] for t in teachers]
print('Distinct names:', len(set(names)))
print('Duplicate names:', [n for n, c in Counter(names).items() if c > 1][:20])

# Sheet1: two columns
ws1 = wb['Sheet1']
rows1 = list(ws1.iter_rows(values_only=True))
left = [r[1] for r in rows1[2:] if len(r) > 1 and r[0] is not None]
right = [r[7] for r in rows1[2:] if len(r) > 7 and r[6] is not None]
print('Sheet1 left count:', len(left), 'Distinct:', len(set(left)))
print('Sheet1 right count:', len(right), 'Distinct:', len(set(right)))

# Sheet2: full names list
ws2 = wb['Sheet2']
rows2 = list(ws2.iter_rows(values_only=True))
s2_names = [r[0] for r in rows2 if r[0]]
print('Sheet2 count:', len(s2_names), 'Distinct:', len(set(s2_names)))
