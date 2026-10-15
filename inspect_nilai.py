#!/usr/bin/env python3
"""Inspect the Data Nilai Excel file: sheet names and first rows."""
import openpyxl

path = "/opt/data/home/workspace/dev/pondok-banat/.data-pondokanat/Data Nilai Semester 1 TA 2026_2027.xlsx"
wb = openpyxl.load_workbook(path, read_only=True)
print("Sheets:", wb.sheetnames)
for name in wb.sheetnames:
    ws = wb[name]
    print(f"--- {name}: {ws.max_row} rows x {ws.max_column} cols ---")
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i < 4:
            print(row)
        else:
            break
wb.close()
