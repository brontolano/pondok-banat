import openpyxl, os

base = "/opt/data/home/workspace/dev/pondok-banat/.data-pondokbanat"
files = [
    "DATA PARA PENGAJAR BANAT.xlsx",
    "Data Banat 2026-2027.xlsx",
    "Data Nilai Semester 1 TA 2026_2027.xlsx",
]

for f in files:
    path = os.path.join(base, f)
    print("=" * 60)
    print(path)
    wb = openpyxl.load_workbook(path, read_only=True)
    for ws in wb.worksheets:
        print("  sheet:", ws.title, "dims:", ws.max_row, "x", ws.max_column)
    wb.close()
