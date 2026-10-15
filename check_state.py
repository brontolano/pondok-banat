#!/usr/bin/env python3
"""Check DB state and Excel file structure."""
import sqlite3
import openpyxl
import os

base = "/opt/data/home/workspace/dev/pondok-banat"

# DB summary
conn = sqlite3.connect(os.path.join(base, "dev.db"))
cur = conn.cursor()
print("=== DATABASE SUMMARY ===")
print(f"users: {cur.execute('SELECT COUNT(*) FROM users').fetchone()[0]}")
print(f"roles: {cur.execute('SELECT COUNT(*) FROM roles').fetchone()[0]}")
print(f"role_permissions: {cur.execute('SELECT COUNT(*) FROM role_permissions').fetchone()[0]}")
print(f"students: {cur.execute('SELECT COUNT(*) FROM students').fetchone()[0]}")
print(f"classes: {cur.execute('SELECT COUNT(*) FROM classes').fetchone()[0]}")
print(f"teachers: {cur.execute('SELECT COUNT(*) FROM teachers').fetchone()[0]}")
print(f"subjects: {cur.execute('SELECT COUNT(*) FROM subjects').fetchone()[0]}")
print(f"silabus: {cur.execute('SELECT COUNT(*) FROM silabus').fetchone()[0]}")
print(f"rpp: {cur.execute('SELECT COUNT(*) FROM rpp').fetchone()[0]}")
print(f"absensi: {cur.execute('SELECT COUNT(*) FROM absensi').fetchone()[0]}")
print(f"grades: {cur.execute('SELECT COUNT(*) FROM grades').fetchone()[0]}")
print(f"rapor: {cur.execute('SELECT COUNT(*) FROM rapor').fetchone()[0]}")
print(f"announcements: {cur.execute('SELECT COUNT(*) FROM announcements').fetchone()[0]}")
print(f"messages: {cur.execute('SELECT COUNT(*) FROM messages').fetchone()[0]}")
print(f"notifications: {cur.execute('SELECT COUNT(*) FROM notifications').fetchone()[0]}")

# Students detail
print("\n=== Students ===")
for row in cur.execute('SELECT id, nisn, name, status, kelas_id FROM students ORDER BY nisn'):
    print(dict(row))

# Classes
print("\n=== Classes ===")
for row in cur.execute('SELECT kode, nama, seksi, jenjang, program FROM classes'):
    print(dict(row))

# Teachers sample
print("\n=== Teachers (first 5) ===")
for row in cur.execute('SELECT id, name, jabatan, status, kelas_id FROM teachers LIMIT 5'):
    print(dict(row))

conn.close()

# Excel file inspection
print("\n=== Excel: Data Nilai Semester 1 TA 2026_2027.xlsx ===")
wb = openpyxl.load_workbook(os.path.join(base, ".data-pondokanat", "Data Nilai Semester 1 TA 2026_2027.xlsx"), read_only=True)
print("Sheets:", wb.sheetnames)
for name in wb.sheetnames:
    ws = wb[name]
    rows = list(ws.iter_rows(values_only=True))
    print(f"  {name}: {len(rows)} rows, {ws.max_column} cols")
    # Count data rows (row 2 onwards where col0 is numeric)
    data_cnt = sum(1 for r in rows[1:] if r[0] is not None and isinstance(r[0], (int, float)))
    print(f"    data rows: {data_cnt}")
wb.close()

# Data Banat Excel
print("\n=== Excel: Data Banat 2026-2027.xlsx ===")
wb2 = openpyxl.load_workbook(os.path.join(base, ".data-pondokanat", "Data Banat 2026-2027.xlsx"), read_only=True)
print("Sheets:", wb2.sheetnames)
for name in wb2.sheetnames:
    ws = wb2[name]
    rows = list(ws.iter_rows(values_only=True))
    print(f"  {name}: {len(rows)} rows x {ws.max_column} cols")
wb2.close()
