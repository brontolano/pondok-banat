import openpyxl
from collections import Counter

# Data Nilai sheets
f3 = '.data-pondokbanat/Data Nilai Semester 1 TA 2026_2027.xlsx'
wb3 = openpyxl.load_workbook(f3, read_only=True)
print('Data Nilai sheets:', wb3.sheetnames)
for sn in wb3.sheetnames:
    ws = wb3[sn]
    rows = list(ws.iter_rows(values_only=True))
    print(f'  {sn}: {len(rows)} rows')

# Mata Pelajaran - clean subject list
f2 = '.data-pondokbanat/Data Banat 2026-2027.xlsx'
wb2 = openpyxl.load_workbook(f2, read_only=True)
ws = wb2['Mata Pelajaran']
rows = list(ws.iter_rows(values_only=True))
subjects = set()
for r in rows[2:]:
    if len(r) > 1 and r[1]:
        subjects.add(str(r[1]).strip())
subjects.discard('MATA PELAJARAN')
print('Subjects (clean):')
for s in sorted(subjects):
    print('  ', s)
print('Count:', len(subjects))

# Data Kelas - actual classes
ws2 = wb2['Data Kelas']
rows2 = list(ws2.iter_rows(values_only=True))
# row 1 is header: TAKHOSUS, TAMHIDI A, TAMHIDI B, 1A, 1B, ...
# columns are: Takhosus | (empty) | Tamhidi A | Tamhidi B | 1A | 1B | 2A | 2B | ...
# Let me find distinct section labels in columns 3-9
sections = set()
for r in rows2[2:]:
    for i in range(3, min(len(r), 10)):
        v = r[i]
        if v:
            sections.add(str(v).strip())
print('Section labels:', sorted(sections))
print('Section count:', len(sections))
