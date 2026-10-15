import openpyxl
import os

base = "/opt/data/home/workspace/dev/pondok-banat"
data_dir = os.path.join(base, ".data-pondokbanat")

print("=== Source Excel files ===")
for f in sorted(os.listdir(data_dir)):
    fp = os.path.join(data_dir, f)
    if f.endswith(".xlsx"):
        wb = openpyxl.load_workbook(fp, read_only=True)
        print("\n--- " + f + " ---")
        for sheet in wb.sheetnames:
            ws = wb[sheet]
            print("  Sheet '" + sheet + "': " + str(ws.max_row) + " rows x " + str(ws.max_column) + " cols")
        wb.close()
