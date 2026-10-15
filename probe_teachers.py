#!/usr/bin/env python3
import openpyxl, os, re

base = "/opt/data/home/workspace/dev/pondok-banat/.data-pondokanat" if False else "/opt/data/home/workspace/dev/pondok-banat/.data-pondokbanat"

print("=== DATA PARA PENGAJAR BANAT ===")
wb = openpyxl.load_workbook(os.path.join(base, "DATA PARA PENGAJAR BANAT.xlsx"), read_only=True)
print("sheets:", wb.sheetnames)
for s in wb.sheetnames:
    ws = wb[s]
    rows = list(ws.iter_rows(values_only=True))
    print(f"--- sheet {s!r}: {len(rows)} rows ---")
    for r in rows[:8]:
        print(r)
    if len(rows) > 8:
        print("...")
        for r in rows[-3:]:
            print(r)
wb.close()
