#!/usr/bin/env python3
"""List contents of the .data-pondokanat directory and inspect the Data Nilai Excel file."""
import os
import openpyxl

base = "/opt/data/home/workspace/dev/pondok-banat"
data_dir = None
for e in os.listdir(base):
    if e.startswith(".data"):
        data_dir = os.path.join(base, e)
        print("Found .data dir:", repr(e))

if data_dir is None:
    print("NO .data dir found!")
    raise SystemExit(1)

print("Contents of", data_dir, ":")
for entry in os.listdir(data_dir):
    print("  ", entry)

# Now inspect the Data Nilai Excel file
nilai_file = os.path.join(data_dir, "Data Nilai Semester 1 TA 2026_2027.xlsx")
print("\nInspecting", nilai_file)
if not os.path.exists(nilai_file):
    print("FILE NOT FOUND!")
    raise SystemExit(1)

wb = openpyxl.load_workbook(nilai_file, read_only=True)
print("Sheets:", wb.sheetnames)
for name in wb.sheetnames:
    ws = wb[name]
    print(f"--- {name}: {ws.max_row} rows x {ws.max_column} cols ---")
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i < 4:
            print(row)
        else:
            break
wb.close()
