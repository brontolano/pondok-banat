import openpyxl, os

base = '/opt/data/home/workspace/dev/pondok-banat/.data-pondokanat/'

print('=== Data Banat 2026-2027 ===')
wb = openpyxl.load_workbook(base + 'Data Banat 2026-2027.xlsx', read_only=True)
for sh in wb.sheetnames:
    ws = wb[sh]
    rows = list(ws.iter_rows(values_only=True))
    print('Sheet: %s -> %d rows' % (sh, len(rows)))
    for r in rows[:3]:
        print('   ', r)

print()
print('=== DATA PARA PENGAJAR BANAT ===')
wb2 = openpyxl.load_workbook(base + 'DATA PARA PENGAJAR BANAT.xlsx', read_only=True)
for sh in wb2.sheetnames:
    ws = wb2[sh]
    rows = list(ws.iter_rows(values_only=True))
    print('Sheet: %s -> %d rows' % (sh, len(rows)))
    for r in rows[:2]:
        print('   ', r)

print()
print('=== Data Nilai Semester 1 TA 2026_2027 ===')
wb3 = openpyxl.load_workbook(base + 'Data Nilai Semester 1 TA 2026_2027.xlsx', read_only=True)
for sh in wb3.sheetnames:
    ws = wb3[sh]
    rows = list(ws.iter_rows(values_only=True))
    print('Sheet: %s -> %d rows' % (sh, len(rows)))
    for r in rows[:2]:
        print('   ', r)

print()
print('=== Pdf had mata pelajaran ===')
pdfdir = base + 'Pdf had mata pelajaran/'
files = sorted(os.listdir(pdfdir))
print('Total files: %d' % len(files))
for f in files:
    print('   ', f)
