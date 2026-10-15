import openpyxl, os
from collections import OrderedDict

base = "/opt/data/home/workspace/dev/pondok-banat/.data-pondokbanat"

# === TEACHER COUNT ===
f = os.path.join(base, "DATA PARA PENGAJAR BANAT.xlsx")
wb = openpyxl.load_workbook(f, read_only=True)
ws = wb["Table 1"]
rows = list(ws.iter_rows(values_only=True))

# Check NO sequence in main table (rows 2-88)
main_nos = [r[0] for r in rows[2:89] if r[0] is not None]
print("Main table NO values:", main_nos)
print("Main table NO count:", len(main_nos))
print("Unique main NO:", len(OrderedDict.fromkeys(main_nos)))

# Badal
badal_nos = [r[0] for r in rows[92:99] if r[0] is not None]
print("Badal NO values:", badal_nos)
print("Badal count:", len(badal_nos))

# Unique names in main (excluding header rows)
main_names = []
for r in rows[2:89]:
    no = r[0]
    nama = r[1]
    if no is not None and nama and str(nama).strip() not in ("NAMA PENGAJAR", "DATA BADAL DARI KELAS 8"):
        main_names.append(str(nama).strip())
print("Main unique teacher names:", len(OrderedDict.fromkeys(main_names)))
print("Main count (incl dup):", len(main_names))

badal_names = []
for r in rows[92:99]:
    no = r[0]
    nama = r[1]
    if no is not None and nama:
        badal_names.append(str(nama).strip())
print("Badal unique names:", len(OrderedDict.fromkeys(badal_names)))

total = len(OrderedDict.fromkeys(main_names)) + len(OrderedDict.fromkeys(badal_names))
print("TOTAL unique teachers:", total)

# === SUBJECTS ===
wb2 = openpyxl.load_workbook(os.path.join(base, "Data Banat 2026-2027.xlsx"), read_only=True)
ws2 = wb2["Mata Pelajaran"]
mrows = list(ws2.iter_rows(values_only=True))
struct = {"NO", "MATA PELAJARAN", "Pengajar", "JP", "KELAS 1", "KELAS 2", "KELAS 3",
          "KELAS 4", "KELAS 5", "KELAS 6", "TAKHOSUS", "TAMHIDI", "", "None", "Kelas"}
subjects = OrderedDict()
for r in mrows:
    for c in r:
        if c is None:
            continue
        s = str(c).strip()
        if s in struct or s.isdigit():
            continue
        subjects[s] = True
print("\n=== SUBJECTS (unique) ===")
print("Count:", len(subjects))
for s in sorted(subjects):
    print("   -", s)

# Subjects that have RPP (from PDF list)
rpp_subjects = ["Fiqih","Hadits","Imla","Mahfudhot","Mantiq","Nahwu","Siroh","Sorof",
                "Tafsir","Tajwid","Tasawuf","Tauhid","Ulumul Hadits","Ulumul Qur'an","Usul Fiqih","Balagoh","Faroidh"]
rpp_set = set(rpp_subjects)
all_subjects = set(subjects.keys())
print("\nSubjects with RPP:", len(rpp_set & all_subjects))
print("Subjects in data but not RPP:", sorted(all_subjects - rpp_set))
print("RPP subjects not in data:", sorted(rpp_set - all_subjects))

# === KELAS / SESI ===
wb3 = openpyxl.load_workbook(os.path.join(base, "Data Nilai Semester 1 TA 2026_2027.xlsx"), read_only=True)
ws3 = wb3.worksheets
print("\nNilai sheets:", len(ws3))
print("Sheets:", [w.title.strip() for w in ws3])

wb.close()
wb2.close()
wb3.close()
