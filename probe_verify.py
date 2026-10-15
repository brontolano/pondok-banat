import openpyxl
from collections import Counter, OrderedDict

base = '.data-pondokbanat'

# Subject count
wb2 = openpyxl.load_workbook(base + '/Data Banat 2026-2027.xlsx', read_only=True)
ws2 = wb2['Mata Pelajaran']
mrows = list(ws2.iter_rows(values_only=True))
subjects = set()
for r in mrows[1:]:
    if len(r) > 1 and r[1]:
        s = str(r[1]).strip()
        if s and s != 'MATA PELAJARAN':
            subjects.add(s)
print('Subjects:', len(subjects))
for s in sorted(subjects):
    print('  -', s)
wb2.close()

# Grade sheets
wb3 = openpyxl.load_workbook(base + '/Data Nilai Semester 1 TA 2026_2027.xlsx', read_only=True)
print('\nNilai sheets:', len(wb3.sheetnames))
for w in wb3.sheetnames:
    print('  -', w)
wb3.close()
