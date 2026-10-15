import openpyxl, os

base = "/opt/data/home/workspace/dev/pondok-banat/.data-pondokbanat"

# 1. Teachers/staff
f = os.path.join(base, "DATA PARA PENGAJAR BANAT.xlsx")
wb = openpyxl.load_workbook(f, read_only=True)
ws = wb["Table 1"]
rows = list(ws.iter_rows(values_only=True))
print("Table 1 rows (incl header):", len(rows))
names = [r for r in rows[1:] if r[0]]
print("Data rows:", len(names))

# 2. Mata Pelajaran sheet
f2 = os.path.join(base, "Data Banat 2026-2027.xlsx")
wb2 = openpyxl.load_workbook(f2, read_only=True)
ws2 = wb2["Mata Pelajaran"]
mrows = list(ws2.iter_rows(values_only=True))
print("Mata Pelajaran rows:", len(mrows))
for r in mrows[:20]:
    print("   ", r)
