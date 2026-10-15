#!/usr/bin/env python3
import openpyxl, os

base = "/opt/data/home/workspace/dev/pondok-banat/.data-pondokbanat"

print("=== NILAI SHEETS ===")
wb = openpyxl.load_workbook(os.path.join(base, "Data Nilai Semester 1 TA 2026_2027.xlsx"), read_only=True)
for s in wb.sheetnames:
    ws = wb[s]
    print(repr(s), "rows", ws.max_row, "cols", ws.max_column)
wb.close()

print()
print("=== DATA BANAT 2026-2027 ===")
wb2 = openpyxl.load_workbook(os.path.join(base, "Data Banat 2026-2027.xlsx"), read_only=True)
print("sheets:", wb2.sheetnames)
for s in wb2.sheetnames:
    ws = wb2[s]
    print(repr(s), "rows", ws.max_row, "cols", ws.max_column)
wb2.close()
