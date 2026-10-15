#!/usr/bin/env python3
"""Inspect Excel file structure."""
import openpyxl
import os

base = "/opt/data/home/workspace/dev/pondok-banat"
data_dir = os.path.join(base, ".data-pondokanat")

# Data Nilai
print("=== Data Nilai Semester 1 TA 2026_2027.xlsx ===")
wb = openpyxl.load_workbook(os.path.join(data_dir, "Data Nilai Semester 1 TA 2026_2027.xlsx"), read_only=True)
print("Sheets:", wb.sheetnames)
for name in wb.sheetnames:
    ws = wb[name]
    rows = list(ws.iter_rows(values_only=True))
    data_cnt = sum(1 for r in rows[1:] if r[0] is not None and isinstance(r[0], (int, float)))
    print(f"  {name}: {len(rows)} rows, {ws.max_column} cols, data rows: {data_cnt}")
    if data_cnt > 0:
        print(f"    Sample header: {rows[1][:8]}")
wb.close()

# Data Banat
print("\n=== Data Banat 2026-2027.xlsx ===")
wb2 = openpyxl.load_workbook(os.path.join(data_dir, "Data Banat 2026-2027.xlsx"), read_only=True)
print("Sheets:", wb2.sheetnames)
for name in wb2.sheetnames:
    ws = wb2[name]
    rows = list(ws.iter_rows(values_only=True))
    print(f"  {name}: {len(rows)} rows x {ws.max_column} cols")
wb2.close()

# DATA PARA PENGAJAR BANAT
print("\n=== DATA PARA PENGAJAR BANAT.xlsx ===")
wb3 = openpyxl.load_workbook(os.path.join(data_dir, "DATA PARA PENGAJAR BANAT.xlsx"), read_only=True)
print("Sheets:", wb3.sheetnames)
for name in wb3.sheetnames[:5]:
    ws = wb3[name]
    rows = list(ws.iter_rows(values_only=True))
    print(f"  {name}: {len(rows)} rows x {ws.max_column} cols")
wb3.close()
