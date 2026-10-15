import openpyxl
wb = openpyxl.load_workbook('.data-pondokbanat/Data Nilai Semester 1 TA 2026_2027.xlsx', read_only=True)
print('Sheets:', wb.sheetnames)
for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    print(f'\n=== {sheet_name} ===')
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i < 5:
            print(row)
        else:
            break
    print(f'Total rows: {ws.max_row}')
wb.close()