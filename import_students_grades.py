#!/usr/bin/env python3
"""
Import students and grades from Excel files into dev.db
"""

import sqlite3
import openpyxl
from pathlib import Path

DB_PATH = Path("dev.db")
EXCEL_NILAI = Path(".data-pondokbanat/Data Nilai Semester 1 TA 2026_2027.xlsx")

# Sheet name to class mapping
SHEET_TO_CLASS = {
    "TK": "TK",
    "TA A": "TA A",
    "TB": "TB",
    "1A": "1A",
    "1B": "1B",
    "2A": "2A",
    "2B": "2B",
    "3A": "3A",
    "3B": "3B",
    "4A": "4A",
    "4B": "4B",
    "5A": "5A",
    "5B": "5B",
    "6A": "6A",
    "6B": "6B",
}

def get_db_connection():
    return sqlite3.connect(DB_PATH)

def get_class_id(conn, class_name):
    """Get class ID by name (kode or nama), create if not exists"""
    cursor = conn.cursor()
    cursor.execute("SELECT id FROM classes WHERE kode = ? OR nama = ?", (class_name, class_name))
    result = cursor.fetchone()
    if result:
        return result[0]
    # Create class - need to generate an ID
    import uuid
    new_id = str(uuid.uuid4())
    cursor.execute("INSERT INTO classes (id, kode, nama, seksi, jenjang, program) VALUES (?, ?, ?, ?, ?, ?)", 
                   (new_id, class_name, class_name, '', '', ''))
    conn.commit()
    return new_id

def get_subject_id(conn, subject_name):
    """Get subject ID by name, create if not exists"""
    cursor = conn.cursor()
    cursor.execute("SELECT id FROM subjects WHERE name = ?", (subject_name,))
    result = cursor.fetchone()
    if result:
        return result[0]
    import uuid
    new_id = str(uuid.uuid4())
    cursor.execute("INSERT INTO subjects (id, kode, name) VALUES (?, ?, ?)", 
                   (new_id, subject_name[:10], subject_name))
    conn.commit()
    return new_id

def import_students_and_grades():
    """Main import function"""
    conn = get_db_connection()
    cursor = conn.cursor()
    
    # Clear existing students and grades (keep admin, roles, teachers, classes)
    print("Clearing existing student and grade data...")
    cursor.execute("DELETE FROM grades")
    cursor.execute("DELETE FROM students")
    cursor.execute("DELETE FROM absensi")
    cursor.execute("DELETE FROM rapor")
    conn.commit()
    
    # Parse Excel
    wb = openpyxl.load_workbook(EXCEL_NILAI, read_only=True)
    
    students_imported = 0
    grades_imported = 0
    
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        class_name = SHEET_TO_CLASS.get(sheet_name.strip())
        if not class_name:
            continue
        
        class_id = get_class_id(conn, class_name)
        print(f"\nProcessing {sheet_name} -> class {class_name} (id={class_id[:8]}...)")
        
        rows = list(ws.iter_rows(values_only=True))
        if len(rows) < 4:
            continue
        
        subject_row = rows[1] if len(rows) > 1 else []
        subheader_row = rows[2] if len(rows) > 2 else []
        
        # Parse subject names from row 1
        subjects = []
        for i, val in enumerate(subject_row):
            if val and isinstance(val, str) and val.strip():
                subjects.append((i, val.strip()))
        
        print(f"  Subjects: {[s[1] for s in subjects]}")
        
        # For each data row
        for row_idx in range(3, len(rows)):
            row = rows[row_idx]
            if not row or all(v is None for v in row):
                continue
            
            no = row[0]
            nis = row[1]
            name = row[2]
            
            if not no or not name or (isinstance(name, str) and not name.strip()):
                continue
            
            # Clean NIS
            nis_str = str(int(nis)) if nis and isinstance(nis, (int, float)) else str(nis).strip() if nis else ""
            if not nis_str or nis_str == "None":
                nis_str = f"{class_name}_{int(no)}" if no else f"{class_name}_{row_idx}"
            
            # Insert student
            import uuid
            student_id = str(uuid.uuid4())
            cursor.execute("""
                INSERT OR IGNORE INTO students (id, nisn, name, jk, ttl, kelas_id, status, created_at, updated_at)
                VALUES (?, ?, ?, ?, ?, ?, ?, datetime('now'), datetime('now'))
            """, (student_id, nis_str, name.strip() if isinstance(name, str) else str(name), 'P', None, class_id, 'inactive'))
            
            if cursor.rowcount > 0:
                students_imported += 1
            
            # Get student ID (in case it already existed)
            cursor.execute("SELECT id FROM students WHERE nisn = ?", (nis_str,))
            student_result = cursor.fetchone()
            if not student_result:
                continue
            student_id = student_result[0]
            
            # Parse grades
            col = 3  # Start after NO, NIS, NAME
            for subj_idx, subj_name in subjects:
                subject_id = get_subject_id(conn, subj_name)
                
                # Determine column layout based on class
                if class_name in ["TK", "TA A", "TB"]:
                    # SMT, ABSEN, HASIL (3 columns)
                    if col + 2 < len(row):
                        smt = row[col]
                        absen = row[col + 1]
                        hasil = row[col + 2]
                        
                        if any(v is not None for v in [smt, absen, hasil]):
                            cursor.execute("""
                                INSERT INTO grades (student_id, kelas_id, mapel_id, semester, minggu, mei, smt, absen, hasil, created_at, updated_at)
                                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, datetime('now'), datetime('now'))
                            """, (student_id, class_id, subject_id, '1', 0, None, smt, absen, hasil))
                            grades_imported += 1
                    col += 3
                else:
                    # 1A-6B: MEI, JUNI, JULI, SMT, ABSEN, HASIL (6 columns)
                    if col + 5 < len(row):
                        mei = row[col]
                        juni = row[col + 1]
                        juli = row[col + 2]
                        smt = row[col + 3]
                        absen = row[col + 4]
                        hasil = row[col + 5]
                        
                        # Calculate MEI average from MEI, JUNI, JULI
                        mei_vals = [v for v in [mei, juni, juli] if isinstance(v, (int, float))]
                        mei_avg = sum(mei_vals) / len(mei_vals) if mei_vals else None
                        
                        if any(v is not None for v in [mei_avg, smt, absen, hasil]):
                            cursor.execute("""
                                INSERT INTO grades (student_id, kelas_id, mapel_id, semester, minggu, mei, smt, absen, hasil, created_at, updated_at)
                                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, datetime('now'), datetime('now'))
                            """, (student_id, class_id, subject_id, '1', 0, mei_avg, smt, absen, hasil))
                            grades_imported += 1
                    col += 6
        
        conn.commit()
        print(f"  Class {class_name}: imported")
    
    wb.close()
    conn.commit()
    conn.close()
    
    print(f"\n=== Import Complete ===")
    print(f"Students imported: {students_imported}")
    print(f"Grades imported: {grades_imported}")

if __name__ == "__main__":
    import_students_and_grades()