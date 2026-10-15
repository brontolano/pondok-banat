#!/usr/bin/env python3
import openpyxl, os

base = "/opt/data/home/workspace/dev/pondok-banat/.data-pondokbanat"

print("=== TAHUN 2026-2027 sheet: teacher rows ===")
wb = openpyxl.load_workbook(os.path.join(base, "DATA PARA PENGAJAR BANAT.xlsx"), read_only=True)
ws = wb["TAHUN 2026-2027"]
rows = list(ws.iter_rows(values_only=True))
# header row index 0
data = [r for r in rows[1:] if r[1] and str(r[1]).strip()]
print("data rows in TAHUN sheet:", len(data))
for r in data:
    print(r[0], '|', r[1], '|', r[2], '|', r[3], '|', r[4], '| wali:', r[5])
wb.close()

print()
print("=== Table 1 count ===")
wb2 = openpyxl.load_workbook(os.path.join(base, "DATA PARA PENGAJAR BANAT.xlsx"), read_only=True)
ws2 = wb2["Table 1"]
rows2 = list(ws2.iter_rows(values_only=True))
data2 = [r for r in rows2[1:] if r[1] and str(r[1]).strip()]
print("data rows in Table1:", len(data2))
wb2.close()

# count unique names
wb3 = openpyxl.load_workbook(os.path.join(base, "DATA PARA PENGAJAR BANAT.xlsx"), read_only=True)
ws3 = wb3["Table 1"]
names = set()
for r in ws3.iter_rows(values_only=True):
    if r[1] and str(r[1]).strip():
        names.add(str(r[1]).strip())
print("unique names in Table1:", len(names))
wb3.close()
