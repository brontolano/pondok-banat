import os
import openpyxl

d = [e for e in os.listdir("/opt/data/home/workspace/dev/pondok-banat") if e.startswith(".data")][0]
p = os.path.join("/opt/data/home/workspace/dev/pondok-banat", d, "Data Banat 2026-2027.xlsx")
wb = openpyxl.load_workbook(p, read_only=True)
ws = wb["Data Kelas"]
print("Data Kelas max_column:", ws.max_column, "max_row:", ws.max_row)
rows = list(ws.iter_rows(values_only=True))
for i, r in enumerate(rows[:6]):
    print("row", i, ":", r)
wb.close()
