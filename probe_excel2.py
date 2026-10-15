#!/usr/bin/env python3
import openpyxl, os

base = "/opt/data/home/workspace/dev/pondok-banat"
data_dir = os.path.join(base, ".data-pondokanat")

for f in ["DATA PARA PENGAJAR BANAT.xlsx", "Data Banat 2026-2027.xlsx", "Data Nilai Semester 1 TA 2026_2027.xlsx"]:
    p = os.path.join(data_dir, f)
    print("===", f, "===")
    wb = openpyxl.load_workbook(p, read_only=True)
    for sh in wb.sheetnames:
        ws = wb[sh]
        rows = list(ws.iter_rows(values_only=True))
        print("  sheet '%s': %d rows" % (sh, len(rows)))
        for r in rows[:2]:
            print("    ", r)
    print()
