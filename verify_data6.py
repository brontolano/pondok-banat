import openpyxl, os
from collections import OrderedDict

base = "/opt/data/home/workspace/dev/pondok-banat/.data-pondokbanat"

# Teachers - main + badal
f = os.path.join(base, "DATA PARA PENGAJAR BANAT.xlsx")
wb = openpyxl.load_workbook(f, read_only=True)
ws = wb["Table 1"]
rows = list(ws.iter_rows(values_only=True))

# Main table rows 2-88, badal rows 92-98
main_teacher_rows = rows[2:89]   # rows 2..88
badal_rows = rows[92:99]         # rows 92..98

def names_from(rows_slice):
    names = []
    nos = []
    for r in rows_slice:
        no = r[0]
        nama = r[1]
        if no is not None and nama:
            nos.append(no)
            names.append(str(nama).strip())
    return names, nos

mn, mnos = names_from(main_teacher_rows)
bn, bnos = names_from(badal_rows)
print("Main teachers (NO 1-52):", len(mn), "unique NO:", len(OrderedDict.fromkeys(mnos)))
print("Badal teachers:", len(bn), "unique NO:", len(OrderedDict.fromkeys(bnos)))
print("Total data rows:", len(mn)+len(bn))
print("All teacher names:")
for n in mn:
    print("   -", n)
print("Badal names:")
for n in bn:
    print("   -", n)

# Check for duplicates in main
dup = [n for n, c in OrderedDict((n, mn.count(n)) for n in mn).items() if c > 1]
print("Duplicate names:", dup)

# Subjects - cleaned
wb2 = openpyxl.load_workbook(os.path.join(base, "Data Banat 2026-2027.xlsx"), read_only=True)
ws2 = wb2["Mata Pelajaran"]
mrows = list(ws2.iter_rows(values_only=True))
# structural tokens
struct = {"NO", "MATA PELAJARAN", "Pengajar", "JP", "KELAS 1", "KELAS 2", "KELAS 3",
          "KELAS 4", "KELAS 5", "KELAS 6", "TAKHOSUS", "TAMHIDI", "", "None", "None"}
subjects = OrderedDict()
for r in mrows:
    for c in r:
        if c is None:
            continue
        s = str(c).strip()
        if s in struct or s.isdigit():
            continue
        subjects[s] = True
print("\nUnique subject names:", len(subjects))
for s in sorted(subjects):
    print("   -", s)

# Count classes/seksi in nilai sheets
wb3 = openpyxl.load_workbook(os.path.join(base, "Data Nilai Semester 1 TA 2026_2027.xlsx"), read_only=True)
ws3 = wb3.worksheets
print("\nNilai sheets:", len(ws3))
for w in ws3:
    print("   -", w.title)
wb3.close()
wb.close()
