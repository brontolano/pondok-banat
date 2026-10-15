import openpyxl
BASE = '/opt/data/home/workspace/dev/pondok-banat/.data-pondokanat'
wb = openpyxl.load_workbook(BASE + '/Data Nilai Semester 1 TA 2026_2027.xlsx', read_only=True)
print('Sheets:', wb.sheetnames)
for s in wb.sheetnames:
    ws = wb[s]
    print(f'--- {s} (rows={ws.max_row}, cols={ws.max_column}) ---')
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i < 5:
            print(row)
        else:
            break
wb.close()
