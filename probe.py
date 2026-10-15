import openpyxl
from collections import Counter

f2 = '.data-pondokbanat/Data Banat 2026-2027.xlsx'
wb2 = openpyxl.load_workbook(f2, read_only=True)

ws = wb2['Mata Pelajaran']
rows = list(ws.iter_rows(values_only=True))
print('=== Mata Pelajaran rows:', len(rows))
subjects = set()
classes = set()
for r in rows[2:]:
    if len(r) > 1 and r[1]:
        subjects.add(str(r[1]).strip())
    if len(r) > 5 and r[5]:
        classes.add(str(r[5]).strip())
print('Subjects:')
for s in sorted(subjects):
    print('  ', s)
print('Classes:')
for c in sorted(classes):
    print('  ', c)

ws2 = wb2['Data Kelas']
rows2 = list(ws2.iter_rows(values_only=True))
takh = set(); tamh = set()
for r in rows2[2:]:
    for v in r:
        if v:
            s = str(v).strip()
            if 'TAMHIDI' in s or 'TAKHOSUS' in s or 'KELAS' in s:
                takh.add(s)
print('Unique class labels:')
for s in sorted(takh):
    print('  ', s)
