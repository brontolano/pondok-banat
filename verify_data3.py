import openpyxl, os, re
from collections import OrderedDict

base = "/opt/data/home/workspace/dev/pondok-banat/.data-pondokbanat"

# 1. Teachers - Table 1
f = os.path.join(base, "DATA PARA PENGAJAR BANAT.xlsx")
wb = openpyxl.load_workbook(f, read_only=True)
ws = wb["Table 1"]
rows = list(ws.iter_rows(values_only=True))
# Table 1 headers: 7 cols (NO, NAMA, JENJANG?, ...?). Let's inspect header
print("TABLE1 header:", rows[0])
names = []
for r in rows[1:]:
    # find non-empty col
    vals = [str(c).strip() for c in r if c is not None and str(c).strip() != ""]
    if vals:
        names.append(vals)
print("Table1 data rows:", len(names))
print("Sample rows:", names[:5])

# unique names by first col (NO) may be 61; check unique
no_vals = [r[0] for r in rows[1:] if r[0]]
print("Unique NO count:", len(OrderedDict.fromkeys(no_vals)))

# scan other sheets for NAMA column
for sn in ["Sheet1", "Sheet2"]:
    wss = wb[sn]
    srows = list(wss.iter_rows(values_only=True))
    print(sn, "dims", wss.max_row, "header", srows[0] if srows else None)
wb.close()

# 2. Mata Pelajaran unique subjects
wb2 = openpyxl.load_workbook(os.path.join(base, "Data Banat 2026-2027.xlsx"), read_only=True)
ws2 = wb2["Mata Pelajaran"]
mrows = list(ws2.iter_rows(values_only=True))
subjects = OrderedDict()
for r in mrows[2:]:  # skip header rows
    for c in r:
        if c is not None and str(c).strip() != "":
            subjects[str(c).strip()] = True
print("Unique subject names:", len(subjects))
for s in subjects:
    print("   -", s)
wb2.close()
